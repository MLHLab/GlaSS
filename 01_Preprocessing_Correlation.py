# ============================================================
# GLACIER MASS-BALANCE MANUSCRIPT-CONSISTENT PIPELINE
# SPYDER STAGED VERSION
# ============================================================
# This script is one stage of a four-stage workflow.
# Run the stages sequentially. Each stage saves a compact pickle
# checkpoint that is consumed by the next stage.
# ============================================================

import pickle
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
import xgboost

from sklearn.cluster import KMeans
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor, StackingRegressor
from sklearn.linear_model import LinearRegression
from sklearn.model_selection import train_test_split, GridSearchCV, KFold, cross_val_score, cross_validate
from sklearn.preprocessing import StandardScaler, MinMaxScaler
from sklearn.metrics import r2_score, mean_absolute_error, mean_squared_error
from sklearn.inspection import permutation_importance
from sklearn.svm import SVR
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "GlacierDataset_175Vars.csv"
OUTPUT_DIR = BASE_DIR / "Glacier_Results"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

STAGE1_FILE = OUTPUT_DIR / "Stage1_Preprocessed.pkl"
STAGE2_FILE = OUTPUT_DIR / "Stage2_SSL_GridSearch.pkl"
STAGE3_FILE = OUTPUT_DIR / "Stage3_Retained_FinalCV.pkl"
OUTPUT_EXCEL = OUTPUT_DIR / "All Results.xlsx"

# Output files generated during Stage 1
CORR_MATRIX_FILE = OUTPUT_DIR / "correlation_matrix.png"
CORR_REDUCED_FILE = OUTPUT_DIR / "GlacierDataset_170Vars_Correlation_Reduced.csv"

# 2. LOAD ORIGINAL DATA
# ============================================================

print('\n' + '=' * 70)
print('STEP 1: LOADING ORIGINAL DATA')
print('=' * 70)

data = pd.read_csv(INPUT_FILE)

# Safety check for the updated dataset structure
assert data.shape[1] >= 2, (
    'The dataset must contain predictor columns and one target column.'
)

print(
    'Number of columns loaded:',
    data.shape[1]
)

print('Original dataset shape:', data.shape)

print('\nFirst five columns:')
print(data.columns[:5].tolist())

print('\nLast five columns:')
print(data.columns[-5:].tolist())

print('\nTotal missing values:', data.isnull().sum().sum())


# Identifier columns are retained in the source data but are NOT used
# as predictive variables, as specified in the manuscript.
IDENTIFIER_COLUMNS = [
    c for c in ['WGMS_ID', 'YR']
    if c in data.columns
]

target_column = data.columns[-1]

feature_columns = [
    c for c in data.columns[:-1]
    if c not in IDENTIFIER_COLUMNS
]

print('\nIdentifier columns excluded from model training:', IDENTIFIER_COLUMNS)
print('Number of predictor variables used for modelling:', len(feature_columns))
print('Target variable:', target_column)

print('\nPredictor columns used for modelling:')
print(feature_columns)

print('\nTarget summary:')
print(data[target_column].describe())

# ============================================================
# 3. DATA GAP-FILLING + CORRELATION-BASED FEATURE REDUCTION
# ============================================================

print('\n' + '=' * 70)
print('STEP 2: GAP-FILLING AND CORRELATION-BASED FEATURE REDUCTION')
print('=' * 70)

# ----------------------------------------------------------------
# 3A. Gap filling
# ----------------------------------------------------------------
# The manuscript specifies linear interpolation between the nearest
# valid observations in a predictor time series. If WGMS_ID and YR
# are available, interpolation is performed separately for each
# glacier and in chronological order. Otherwise, interpolation is
# performed in the row order of the supplied dataset.
# ----------------------------------------------------------------

data_work = data.copy()

missing_before = int(
    data_work[feature_columns].isna().sum().sum()
)

print('Missing predictor values before interpolation:', missing_before)

if missing_before > 0:

    if 'WGMS_ID' in data_work.columns and 'YR' in data_work.columns:

        data_work = data_work.sort_values(
            ['WGMS_ID', 'YR']
        ).reset_index(drop=True)

        for col in feature_columns:
            data_work[col] = (
                data_work
                .groupby('WGMS_ID', sort=False)[col]
                .transform(
                    lambda s: s.interpolate(
                        method='linear',
                        limit_area='inside'
                    )
                )
            )

    else:

        for col in feature_columns:
            data_work[col] = data_work[col].interpolate(
                method='linear',
                limit_area='inside'
            )

missing_after = int(
    data_work[feature_columns].isna().sum().sum()
)

print('Missing predictor values after interpolation:', missing_after)

if missing_after > 0:
    remaining = (
        data_work[feature_columns]
        .isna()
        .sum()
        .loc[lambda s: s > 0]
        .sort_values(ascending=False)
    )

    print('\nWARNING: Some missing values remain after interpolation.')
    print(remaining)

    raise ValueError(
        'Missing predictor values remain after the manuscript-specified '
        'linear interpolation. Check missing values at the beginning/end '
        'of glacier time series before continuing.'
    )

# Keep identifiers available for traceability, but never include them
# in model training.
data = data_work.copy()
data_full = data.copy()

data1 = data[feature_columns].copy()

print('Predictor matrix after gap filling:', data1.shape)

# Target
y_original = data[target_column].copy()

print('Target shape:', y_original.shape)

# ----------------------------------------------------------------
# 3B. Pearson correlation matrix
# ----------------------------------------------------------------

c2 = data1.corr(method='pearson')

print('Correlation matrix shape:', c2.shape)

plt.figure(figsize=(30, 30))
sns.heatmap(
    c2,
    annot=False,
    cmap='coolwarm',
    center=0
)

plt.title('Pearson Correlation Matrix')
plt.savefig(
    CORR_MATRIX_FILE,
    dpi=300,
    bbox_inches='tight'
)
plt.close()

# ----------------------------------------------------------------
# 3C. Identify highly correlated predictor pairs
# ----------------------------------------------------------------

corr_threshold = 0.90

corrs = []

for i in range(len(feature_columns)):

    for j in range(i + 1, len(feature_columns)):

        if abs(c2.iloc[i, j]) >= corr_threshold:
            corrs.append([i, j])

print(
    'Number of highly correlated pairs:',
    len(corrs)
)

# ----------------------------------------------------------------
# 3D. Reject the second variable in each highly correlated pair
# ----------------------------------------------------------------

rej = [
    feature_columns[pair[1]]
    for pair in corrs
]

rej = list(dict.fromkeys(rej))

print(
    'Number of rejected correlated variables:',
    len(rej)
)

print('\nRejected variables:')
print(rej)

# ----------------------------------------------------------------
# 3E. Create reduced dataset
# ----------------------------------------------------------------

data2a = data.drop(
    columns=rej
).copy()

print(
    '\nCorrelation-reduced dataset shape:',
    data2a.shape
)

# Save reduced dataset
data2a.to_csv(
    CORR_REDUCED_FILE,
    index=False
)

print('\nSaved correlation-reduced dataset:')
print(CORR_REDUCED_FILE)


# ============================================================
# 4. LOAD CORRELATION-REDUCED DATA
# ============================================================

print('\n' + '=' * 70)
print('STEP 3: PREPARING REDUCED DATA FOR FEATURE RECONSTRUCTION')
print('=' * 70)

data = pd.read_csv(
    CORR_REDUCED_FILE
)

print(
    'Correlation-reduced dataset shape:',
    data.shape
)

print(
    '\nTotal missing values:',
    data.isnull().sum().sum()
)

data2 = data.copy()

target_column = data2.columns[-1]

feature_columns_reduced = [
    c for c in data2.columns
    if c not in IDENTIFIER_COLUMNS + [target_column]
]

print(
    '\nNumber of retained predictors:',
    len(feature_columns_reduced)
)

print(
    'Target column:',
    target_column
)

print(
    'data2 shape:',
    data2.shape
)

# ============================================================
# 4A. PREPARE FEATURES AND TARGET FOR STAGE 2
# ============================================================
# Preserve the unscaled correlation-reduced predictors for the
# non-SSL ablation and standardize predictors for SSL reconstruction.

X0_df = data2[feature_columns_reduced].copy()
y0_series = data2[target_column].copy()

X00 = X0_df.copy()
y00 = y0_series.copy()

print('X0 before scaling:', X0_df.shape)
print('y0:', y0_series.shape)
print('Number of retained predictors:', X0_df.shape[1])

scaler = StandardScaler()
X0 = scaler.fit_transform(X0_df)
y0 = y0_series.to_numpy(dtype=float)

print('Scaled X0 shape:', X0.shape)
print('y0 shape:', y0.shape)


# ============================================================
# SAVE STAGE 1 CHECKPOINT
# ============================================================
stage1 = {
    "data_full": data_full,
    "data2": data2,
    "IDENTIFIER_COLUMNS": IDENTIFIER_COLUMNS,
    "target_column": target_column,
    "feature_columns": feature_columns,
    "feature_columns_reduced": feature_columns_reduced,
    "X0_df": X0_df,
    "y0_series": y0_series,
    "X00": X00,
    "y00": y00,
    "scaler": scaler,
    "X0": X0,
    "y0": y0,
}
with open(STAGE1_FILE, "wb") as f:
    pickle.dump(stage1, f, protocol=pickle.HIGHEST_PROTOCOL)

print("\nSTAGE 1 COMPLETE")
print("Checkpoint saved:", STAGE1_FILE)
print("Run 02_SSL_and_GridSearch.py next.")
