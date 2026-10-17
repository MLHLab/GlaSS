# ============================================================
# GLACIER MASS-BALANCE MANUSCRIPT-CONSISTENT PIPELINE
# SPYDER STAGED VERSION
# ============================================================
# This script is one stage of a four-stage workflow.
# Run the stages sequentially. Each stage saves a compact pickle
# checkpoint that is consumed by the next stage.
# ============================================================

import pickle
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
import xgboost

from sklearn.cluster import KMeans
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor, StackingRegressor
from sklearn.linear_model import LinearRegression
from sklearn.model_selection import train_test_split, GridSearchCV, KFold, cross_val_score, cross_validate
from sklearn.preprocessing import StandardScaler, MinMaxScaler
from sklearn.metrics import r2_score, mean_absolute_error, mean_squared_error
from sklearn.inspection import permutation_importance
from sklearn.svm import SVR
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "GlacierDataset_170Vars.csv"
OUTPUT_DIR = BASE_DIR / "Glacier_Results"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

STAGE1_FILE = OUTPUT_DIR / "Stage1_Preprocessed.pkl"
STAGE2_FILE = OUTPUT_DIR / "Stage2_SSL_GridSearch.pkl"
STAGE3_FILE = OUTPUT_DIR / "Stage3_Retained_FinalCV.pkl"
OUTPUT_EXCEL = OUTPUT_DIR / "All Results.xlsx"

if not STAGE3_FILE.exists():
    raise FileNotFoundError(f"Stage 3 checkpoint not found: {STAGE3_FILE}\nRun 03_Rejection_and_RetainedData.py first.")
with open(STAGE3_FILE, "rb") as f:
    stage3 = pickle.load(f)

globals().update(stage3)
print("Loaded Stage 3:", STAGE3_FILE)


if "data_full" not in globals():
    raise KeyError(
        "Stage 3 checkpoint does not contain data_full. "
        "Please rerun Stage 1 with the corrected script."
    )

if "data2" not in globals():
    raise KeyError(
        "Stage 3 checkpoint does not contain data2. "
        "Please rerun Stage 1 with the corrected script."
    )

# Original gap-filled dataset used for full-data reporting.
data = data_full.copy()

# Correlation-reduced dataset used for reduced-data reporting.
data2a = data2.copy()

print("Restored original data alias:", data.shape)
print("Restored reduced data alias:", data2a.shape)


if "FINAL_GBR_PARAMS" not in globals():
    if "BEST_GBR_PARAMS" in globals():
        FINAL_GBR_PARAMS = dict(BEST_GBR_PARAMS)
    else:
        FINAL_GBR_PARAMS = {
            "n_estimators": 600,
            "learning_rate": 0.05,
            "max_depth": 6,
            "subsample": 0.8,
        }

if "FINAL_XGB_PARAMS" not in globals():
    if "BEST_XGB_PARAMS" in globals():
        FINAL_XGB_PARAMS = dict(BEST_XGB_PARAMS)
    else:
        FINAL_XGB_PARAMS = {
            "n_estimators": 600,
            "learning_rate": 0.05,
            "max_depth": 4,
            "subsample": 0.8,
            "colsample_bytree": 0.8,
        }

if "BEST_SVR_PARAMS" not in globals():
    # Stage 2 normally supplies this. Keep a safe fallback so that
    # Stage 4 does not fail merely because an older Stage 2 checkpoint
    # was used.
    BEST_SVR_PARAMS = {
        "C": 1.0,
        "epsilon": 0.1,
        "gamma": "scale",
        "kernel": "rbf",
    }


def create_stacking_regressor():
    """
    Create the final GBR + XGBR + Linear Regression stacking model.

    The GBR/XGBR hyperparameters are taken from the completed Stage 2
    GridSearch checkpoint whenever available.
    """

    gbr = GradientBoostingRegressor(
        **FINAL_GBR_PARAMS,
        random_state=42
    )

    xgb = xgboost.XGBRegressor(
        **FINAL_XGB_PARAMS,
        objective="reg:squarederror",
        random_state=42,
        n_jobs=-1
    )

    lr = LinearRegression()

    return StackingRegressor(
        estimators=[
            ("gbr", gbr),
            ("xgb", xgb)
        ],
        final_estimator=lr,
        n_jobs=-1
    )


# File used for the final fold-wise prediction table.
CV_PRED_FILE = OUTPUT_DIR / "Final_CV_Actual_vs_Predicted.csv"

# Standardized correlation-reduced predictors for the ablation study.
# Stage 3 retains X0_merge in the original/reduced feature scale; the
# Stage 1 scaler is reused so that the no-SSL ablation uses the same
# standardized representation as the SSL workflow.
if "X0_merge_standardized" not in globals():
    if "scaler" not in globals():
        raise KeyError(
            "The Stage 3 checkpoint does not contain the Stage 1 scaler. "
            "Please rerun Stage 1 and then Stage 2 and Stage 3."
        )
    X0_merge_standardized = scaler.transform(
        X0_merge.to_numpy(dtype=float)
    )

# Combined retained/rejected tables used only for the final summary.
Xy0_merge = pd.concat(
    [X0_merge.reset_index(drop=True),
     y0_merge.reset_index(drop=True).rename(target_column)],
    axis=1
)

Xy0_merge_r = pd.concat(
    [X0_rejected.reset_index(drop=True),
     y0_rejected.reset_index(drop=True).rename(target_column)],
    axis=1
)
# 13. FINAL 5-FOLD CROSS-VALIDATION
# ============================================================

print('\n' + '=' * 70)
print('STEP 9: FINAL 5-FOLD CROSS-VALIDATION')
print('=' * 70)

kf = KFold(
    n_splits=5,
    shuffle=True,
    random_state=42
)

cv_rows = []
fold_prediction_tables = []

for fold_number, (train_idx, test_idx) in enumerate(
    kf.split(X_merge),
    start=1
):

    print(
        f'\nProcessing final CV fold {fold_number}/5'
    )

    X_train = X_merge[train_idx]
    X_test = X_merge[test_idx]

    y_train = y_merge[train_idx]
    y_test = y_merge[test_idx]

    model_fold = create_stacking_regressor()

    fit_start = __import__('time').perf_counter()

    model_fold.fit(
        X_train,
        y_train
    )

    y_pred_train = model_fold.predict(
        X_train
    )

    y_pred_test = model_fold.predict(
        X_test
    )

    elapsed = __import__('time').perf_counter() - fit_start

    r2 = r2_score(y_test, y_pred_test)
    mae = mean_absolute_error(y_test, y_pred_test)
    rmse = np.sqrt(mean_squared_error(y_test, y_pred_test))

    cv_rows.append({
        'Fold': fold_number,
        'R2': r2,
        'MAE': mae,
        'RMSE': rmse,
        'Time (s)': elapsed
    })

    fold_prediction_tables.append(
        pd.DataFrame({
            'y_actual': y_test,
            'y_predicted': y_pred_test,
            'fold': fold_number
        })
    )

cv_results_df = pd.DataFrame(cv_rows)

print('\nFold-wise CV results:')
print(cv_results_df)

print('\nMean ± SD:')
print(
    'R2   :',
    cv_results_df['R2'].mean(),
    '±',
    cv_results_df['R2'].std()
)

print(
    'MAE  :',
    cv_results_df['MAE'].mean(),
    '±',
    cv_results_df['MAE'].std()
)

print(
    'RMSE :',
    cv_results_df['RMSE'].mean(),
    '±',
    cv_results_df['RMSE'].std()
)

print(
    'Time :',
    cv_results_df['Time (s)'].mean(),
    '±',
    cv_results_df['Time (s)'].std()
)


# ============================================================
# 14. FOLD-WISE ACTUAL VS PREDICTED VALUES
# ============================================================

print('\n' + '=' * 70)
print('STEP 10: FOLD-WISE ACTUAL VS PREDICTED VALUES')
print('=' * 70)

y0_ac_pr_ts_fin = pd.concat(
    fold_prediction_tables,
    axis=0,
    ignore_index=True
)

print(
    '\nFinal actual-vs-predicted table shape:',
    y0_ac_pr_ts_fin.shape
)

y0_ac_pr_ts_fin.to_csv(
    CV_PRED_FILE,
    index=False
)

print(
    'Saved final CV prediction file:',
    CV_PRED_FILE
)


# ============================================================
# 15. EXCEL RESULT MATRIX
# ============================================================
# The workbook reproduces the six sheets in the supplied
# "All Results.xlsx":
#
#   1. Model Performance
#   2. Comp. Study
#   3. Ablation
#   4. Topographical Param.
#   5. Reduced Feat. Imp.
#   6. Reduced Month-wise Feat. Imp.
#
# Every numerical result is generated by the program; the
# workbook is not populated with hard-coded results from the
# reference workbook.
# ============================================================

from time import perf_counter
from sklearn.inspection import permutation_importance


def evaluate_single_model_cv(
    model_factory,
    X,
    y,
    n_splits=5,
    random_state=42
):
    """
    Evaluate a model using explicit K-fold CV.

    Returns:
        metrics_df
        predictions_df
        fitted_models
    """

    kf_local = KFold(
        n_splits=n_splits,
        shuffle=True,
        random_state=random_state
    )

    rows = []
    predictions = []
    fitted_models = []

    for fold, (tr, ts) in enumerate(
        kf_local.split(X),
        start=1
    ):

        model = model_factory()

        t0 = perf_counter()

        model.fit(
            X[tr],
            y[tr]
        )

        elapsed = perf_counter() - t0

        pred = model.predict(
            X[ts]
        )

        rows.append({
            'Fold': fold,
            'R2': r2_score(y[ts], pred),
            'MAE': mean_absolute_error(y[ts], pred),
            'RMSE': np.sqrt(
                mean_squared_error(
                    y[ts],
                    pred
                )
            ),
            'Time (s)': elapsed
        })

        predictions.append(
            pd.DataFrame({
                'y_actual': y[ts],
                'y_predicted': pred,
                'fold': fold
            })
        )

        fitted_models.append(model)

    return (
        pd.DataFrame(rows),
        pd.concat(
            predictions,
            ignore_index=True
        ),
        fitted_models
    )


def summary_metrics(
    metrics_df
):
    """
    Return mean and standard deviation for the
    three performance metrics and runtime.
    """

    return {
        'R2_mean': metrics_df['R2'].mean(),
        'R2_std': metrics_df['R2'].std(),
        'MAE_mean': metrics_df['MAE'].mean(),
        'MAE_std': metrics_df['MAE'].std(),
        'RMSE_mean': metrics_df['RMSE'].mean(),
        'RMSE_std': metrics_df['RMSE'].std(),
        'Time_mean': metrics_df['Time (s)'].mean(),
        'Time_std': metrics_df['Time (s)'].std()
    }


def make_stacking_factory():
    return lambda: create_stacking_regressor()


# ------------------------------------------------------------
# A. Model Performance
# ------------------------------------------------------------
print('\nCreating Model Performance results...')

X_170 = data_full[feature_columns].to_numpy(dtype=float)
y_all = data_full[target_column].to_numpy(dtype=float)

perf_170, _, _ = evaluate_single_model_cv(
    make_stacking_factory(),
    X_170,
    y_all
)

X_reduced_full = data2[feature_columns_reduced].to_numpy(dtype=float)

perf_reduced, _, _ = evaluate_single_model_cv(
    make_stacking_factory(),
    X_reduced_full,
    y_all
)

# 3720 x reduced: final retained dataset
perf_3720 = cv_results_df.copy()

s170 = summary_metrics(perf_170)
sred = summary_metrics(perf_reduced)
s3720 = summary_metrics(perf_3720)

model_performance_rows = [
    [
        f'{len(data)}×{len(feature_columns)}',
        s170['R2_mean'], s170['R2_std'],
        s170['MAE_mean'], s170['MAE_std'],
        s170['RMSE_mean'], s170['RMSE_std'],
        s170['Time_mean'], s170['Time_std']
    ],
    [
        f'{len(data)}×{X_reduced_full.shape[1]}',
        sred['R2_mean'], sred['R2_std'],
        sred['MAE_mean'], sred['MAE_std'],
        sred['RMSE_mean'], sred['RMSE_std'],
        sred['Time_mean'], sred['Time_std']
    ],
    [
        f'{X_merge.shape[0]}×{X_merge.shape[1]}',
        s3720['R2_mean'], s3720['R2_std'],
        s3720['MAE_mean'], s3720['MAE_std'],
        s3720['RMSE_mean'], s3720['RMSE_std'],
        s3720['Time_mean'], s3720['Time_std']
    ]
]

model_performance_df = pd.DataFrame(
    model_performance_rows,
    columns=[
        'Input Dimension',
        'R2', 'R2 SD',
        'MAE', 'MAE SD',
        'RMSE', 'RMSE SD',
        'Time (s)', 'Time SD'
    ]
)


# ------------------------------------------------------------
# B. Comparative Study
# ------------------------------------------------------------
print('\nCreating Comparative Study results...')

# Use final retained reconstructed representation.
comp_svm, _, _ = evaluate_single_model_cv(
    lambda: SVR(**BEST_SVR_PARAMS),
    X_merge,
    y_merge
)

comp_prop = perf_3720.copy()

comp_models = {
    'SVM for Regression': summary_metrics(comp_svm),
    'Proposed Model': summary_metrics(comp_prop)
}

# The manuscript reports CNN-Attention and LSTM-Attention as external
# benchmark architectures, but it does not specify their layer
# configurations, input reshaping, attention mechanism, optimizer,
# regularization, epochs, or other training details. The previous code
# incorrectly substituted MLPRegressor for both models. That is removed
# here rather than presenting an unrelated MLP as a CNN/LSTM benchmark.
#
# To reproduce those two benchmark rows from first principles, their
# exact architecture/training specification must be supplied.
RUN_DEEP_COMPARISON = False

if RUN_DEEP_COMPARISON:
    raise NotImplementedError(
        'CNN-Attention and LSTM-Attention cannot be reproduced '
        'faithfully from the manuscript alone because their exact '
        'architectures and training hyperparameters are not specified.'
    )

comp_rows = []

for model_name in [
    'SVM for Regression',
    'CNN-Attention with Regularization',
    'LSTM-Attention Model',
    'Proposed Model'
]:

    if model_name not in comp_models:
        continue

    s = comp_models[model_name]

    comp_rows.append([
        model_name,
        s['R2_mean'], s['R2_std'],
        s['MAE_mean'], s['MAE_std'],
        s['RMSE_mean'], s['RMSE_std'],
        s['Time_mean'], s['Time_std']
    ])

comp_study_df = pd.DataFrame(
    comp_rows,
    columns=[
        'Model',
        'R2', 'R2 SD',
        'MAE', 'MAE SD',
        'RMSE', 'RMSE SD',
        'Time (S)', 'Time SD'
    ]
)


# ------------------------------------------------------------
# C. Ablation Study
# ------------------------------------------------------------
print('\nCreating Ablation results...')

ablation_rows = []


def add_ablation(
    name,
    factory,
    X,
    y
):

    metrics, _, _ = evaluate_single_model_cv(
        factory,
        X,
        y
    )

    s = summary_metrics(metrics)

    ablation_rows.append([
        name,
        s['R2_mean'], s['R2_std'],
        s['MAE_mean'], s['MAE_std'],
        s['RMSE_mean'], s['RMSE_std'],
        s['Time_mean'], s['Time_std']
    ])


# Without self-supervised learning:
# use the ORIGINAL standardized correlation-reduced predictors
# for exactly the same retained observations.
X_ablation_without_ssl = X0_merge_standardized
y_ablation_without_ssl = y0_merge.to_numpy(dtype=float)

add_ablation(
    'GBR without Self-Supervised Learning',
    lambda: GradientBoostingRegressor(
        **FINAL_GBR_PARAMS,
        random_state=42
    ),
    X_ablation_without_ssl,
    y_ablation_without_ssl
)

add_ablation(
    'XGBR without Self-Supervised Learning',
    lambda: xgboost.XGBRegressor(
        **FINAL_XGB_PARAMS,
        objective='reg:squarederror',
        random_state=42,
        n_jobs=-1
    ),
    X_ablation_without_ssl,
    y_ablation_without_ssl
)

add_ablation(
    'LR without Self-Supervised Learning',
    lambda: LinearRegression(),
    X_ablation_without_ssl,
    y_ablation_without_ssl
)

add_ablation(
    'Stacking Regressor without Self-Supervised Learning',
    make_stacking_factory(),
    X_ablation_without_ssl,
    y_ablation_without_ssl
)

# With self-supervised learning:
# use the reconstructed representation for exactly the same
# retained observations.
add_ablation(
    'GBR with Self-Supervised Learning',
    lambda: GradientBoostingRegressor(
        **FINAL_GBR_PARAMS,
        random_state=42
    ),
    X_merge,
    y_merge
)

add_ablation(
    'XGBR with Self-Supervised Learning',
    lambda: xgboost.XGBRegressor(
        **FINAL_XGB_PARAMS,
        objective='reg:squarederror',
        random_state=42,
        n_jobs=-1
    ),
    X_merge,
    y_merge
)

add_ablation(
    'LR with Self-Supervised Learning',
    lambda: LinearRegression(),
    X_merge,
    y_merge
)

s = summary_metrics(
    perf_3720
)

ablation_rows.append([
    'Proposed Model',
    s['R2_mean'], s['R2_std'],
    s['MAE_mean'], s['MAE_std'],
    s['RMSE_mean'], s['RMSE_std'],
    s['Time_mean'], s['Time_std']
])

ablation_df = pd.DataFrame(
    ablation_rows,
    columns=[
        'Model',
        'R2', 'R2 SD',
        'MAE', 'MAE SD',
        'RMSE', 'RMSE SD',
        'Time (S)', 'Time SD'
    ]
)


# ------------------------------------------------------------
# D. Topographical Parameter Study
# ------------------------------------------------------------
print('\nCreating Topographical Parameter results...')

# The supplied 170-column dataset contains 168 climatological
# variables + EAS + NOR. The reference workbook additionally
# contains a comparison involving five topographic predictors.
#
# Because those five additional predictors are not present in
# the supplied CSV, the code does not invent them.
#
# If they are available, specify their exact column names here.
TOPOGRAPHIC_COLUMNS = []

topo_rows = []

s_base = summary_metrics(
    perf_3720
)

topo_rows.append([
    '168 climatological features and 2 spatial coordinates',
    s_base['R2_mean'], s_base['R2_std'],
    s_base['MAE_mean'], s_base['MAE_std'],
    s_base['RMSE_mean'], s_base['RMSE_std']
])

if len(TOPOGRAPHIC_COLUMNS) == 5:

    missing = [
        c for c in TOPOGRAPHIC_COLUMNS
        if c not in data.columns
    ]

    if missing:
        raise ValueError(
            'The following topographic columns were not found: '
            + str(missing)
        )

    X_topo_df = pd.concat(
        [
            data_full[feature_columns],
            data_full[TOPOGRAPHIC_COLUMNS]
        ],
        axis=1
    )

    # Remove duplicated columns if a supplied topographic
    # variable is already present.
    X_topo_df = X_topo_df.loc[
        :,
        ~X_topo_df.columns.duplicated()
    ]

    X_topo = X_topo_df.to_numpy(dtype=float)

    topo_metrics, _, _ = evaluate_single_model_cv(
        make_stacking_factory(),
        X_topo,
        y_all
    )

    s_topo = summary_metrics(
        topo_metrics
    )

    topo_rows.append([
        '168 climatological features, 5 topographic predictors, 2 spatial coordinates',
        s_topo['R2_mean'], s_topo['R2_std'],
        s_topo['MAE_mean'], s_topo['MAE_std'],
        s_topo['RMSE_mean'], s_topo['RMSE_std']
    ])

else:

    topo_rows.append([
        '168 climatological features, 5 topographic predictors, 2 spatial coordinates',
        'NOT AVAILABLE',
        '',
        'NOT AVAILABLE',
        '',
        'NOT AVAILABLE',
        ''
    ])

topographical_df = pd.DataFrame(
    topo_rows,
    columns=[
        'Parameter Set',
        'R2', 'R2 SD',
        'MAE (mm w.e. a-1)', 'MAE SD',
        'RMSE (mm w.e. a-1)', 'RMSE SD'
    ]
)


# ------------------------------------------------------------
# E. Reduced Feature Importance
# ------------------------------------------------------------
print('\nCreating Reduced Feature Importance results...')

# Train the proposed model on the final retained dataset.
importance_model = create_stacking_regressor()

importance_model.fit(
    X_merge,
    y_merge
)

if len(feature_columns_reduced) != X_merge.shape[1]:
    raise RuntimeError(
        'Feature-name count does not match final model representation: '
        f'{len(feature_columns_reduced)} names vs {X_merge.shape[1]} features.'
    )

# Permutation importance on the complete retained dataset.
perm = permutation_importance(
    importance_model,
    X_merge,
    y_merge,
    scoring='r2',
    n_repeats=10,
    random_state=42,
    n_jobs=-1
)

importance_df = pd.DataFrame({
    'Feature Variable': feature_columns_reduced,
    'Mean Permutation Importance': perm.importances_mean
})

# The manuscript specifies min-max normalization to [0, 1].
importance_df = importance_df.sort_values(
    'Mean Permutation Importance',
    ascending=False
).reset_index(
    drop=True
)

imp_min = importance_df[
    'Mean Permutation Importance'
].min()

imp_max = importance_df[
    'Mean Permutation Importance'
].max()

if imp_max > imp_min:
    importance_df[
        'Normalized Score'
    ] = (
        importance_df[
            'Mean Permutation Importance'
        ] - imp_min
    ) / (imp_max - imp_min)
else:
    importance_df[
        'Normalized Score'
    ] = 0.0


# ------------------------------------------------------------
# F. Reduced Month-wise Feature Importance
# ------------------------------------------------------------
print('\nCreating Reduced Month-wise Feature Importance results...')

month_records = []

for _, row in importance_df.iterrows():

    feature = row[
        'Feature Variable'
    ]

    if '_' in feature:

        parts = feature.rsplit(
            '_',
            1
        )

        variable = parts[0]
        month = parts[1]

    else:

        variable = feature
        month = ''

    month_records.append([
        month,
        variable,
        row['Normalized Score'],
        row['Mean Permutation Importance']
    ])

monthwise_df = pd.DataFrame(
    month_records,
    columns=[
        'Month',
        'Feature Variable',
        'Normalized Score',
        'Mean Permutation Importance'
    ]
)

# Sort month-wise in calendar order while preserving EAS/NOR.
month_order = [
    'JAN', 'FEB', 'MAR', 'APR',
    'MAY', 'JUN', 'JUL', 'AUG',
    'SEP', 'OCT', 'NOV', 'DEC'
]

month_rank = {
    m: i for i, m in enumerate(month_order)
}

monthwise_df['_month_rank'] = (
    monthwise_df['Month']
    .map(month_rank)
    .fillna(999)
)

monthwise_df = monthwise_df.sort_values(
    ['_month_rank', 'Mean Permutation Importance'],
    ascending=[True, False]
).drop(
    columns=['_month_rank']
).reset_index(
    drop=True
)

# Keep the same min-max normalized score produced for the complete
# feature-importance analysis. No second max-only normalization is
# applied, so the month-wise table remains consistent with the
# manuscript's [0, 1] min-max normalization.


# ============================================================
# 16. WRITE EXCEL WORKBOOK
# ============================================================

print('\n' + '=' * 70)
print('STEP 11: WRITING EXCEL RESULT WORKBOOK')
print('=' * 70)


def write_title(
    ws,
    text,
    end_col
):

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_col
    )

    cell = ws.cell(
        row=1,
        column=1
    )

    cell.value = text
    cell.font = Font(
        bold=True,
        size=12
    )

    cell.alignment = Alignment(
        horizontal='center',
        vertical='center',
        wrap_text=True
    )

    ws.row_dimensions[1].height = 35


def style_header(
    ws,
    row
):

    for cell in ws[row]:

        if cell.value is not None:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal='center',
                vertical='center',
                wrap_text=True
            )


def auto_width(
    ws,
    minimum=10,
    maximum=45
):

    for col_cells in ws.columns:

        col_letter = get_column_letter(
            col_cells[0].column
        )

        max_length = 0

        for cell in col_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            col_letter
        ].width = max(
            minimum,
            min(
                max_length + 2,
                maximum
            )
        )


def add_number_format(
    ws,
    start_row,
    end_row,
    start_col=2,
    end_col=None
):

    if end_col is None:
        end_col = ws.max_column

    for row in ws.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_col,
        max_col=end_col
    ):

        for cell in row:

            if isinstance(
                cell.value,
                (float, int)
            ):

                cell.number_format = '0.0000'


with pd.ExcelWriter(
    OUTPUT_EXCEL,
    engine='openpyxl'
) as writer:

    # --------------------------------------------------------
    # Sheet 1: Model Performance
    # --------------------------------------------------------

    ws_name = 'Model Performance'

    # Write title separately to reproduce reference layout.
    model_performance_df.to_excel(
        writer,
        sheet_name=ws_name,
        startrow=2,
        index=False
    )

    ws = writer.book[ws_name]

    write_title(
        ws,
        'Performance with 168 climatological features and 2 spatial coordinates\n'
        'Correlation as feature reduction',
        9
    )

    style_header(
        ws,
        3
    )

    add_number_format(
        ws,
        4,
        ws.max_row
    )

    # --------------------------------------------------------
    # Sheet 2: Comparative Study
    # --------------------------------------------------------

    ws_name = 'Comp. Study'

    comp_study_df.to_excel(
        writer,
        sheet_name=ws_name,
        startrow=2,
        index=False
    )

    ws = writer.book[ws_name]

    write_title(
        ws,
        f'Final retained dataset ({X_merge.shape[0]} X {X_merge.shape[1]})',
        9
    )

    style_header(
        ws,
        3
    )

    add_number_format(
        ws,
        4,
        ws.max_row
    )

    # --------------------------------------------------------
    # Sheet 3: Ablation
    # --------------------------------------------------------

    ws_name = 'Ablation'

    ablation_df.to_excel(
        writer,
        sheet_name=ws_name,
        startrow=2,
        index=False
    )

    ws = writer.book[ws_name]

    write_title(
        ws,
        f'Final retained dataset ({X_merge.shape[0]} X {X_merge.shape[1]})',
        9
    )

    style_header(
        ws,
        3
    )

    add_number_format(
        ws,
        4,
        ws.max_row
    )

    # --------------------------------------------------------
    # Sheet 4: Topographical Param.
    # --------------------------------------------------------

    ws_name = 'Topographical Param.'

    topographical_df.to_excel(
        writer,
        sheet_name=ws_name,
        startrow=1,
        index=False
    )

    ws = writer.book[ws_name]

    style_header(
        ws,
        2
    )

    add_number_format(
        ws,
        3,
        ws.max_row
    )

    # --------------------------------------------------------
    # Sheet 5: Reduced Feat. Imp.
    # --------------------------------------------------------

    ws_name = 'Reduced Feat. Imp.'

    importance_df.to_excel(
        writer,
        sheet_name=ws_name,
        index=False
    )

    ws = writer.book[ws_name]

    style_header(
        ws,
        1
    )

    add_number_format(
        ws,
        2,
        ws.max_row
    )

    # --------------------------------------------------------
    # Sheet 6: Reduced Month-wise Feat. Imp.
    # --------------------------------------------------------

    ws_name = 'Reduced Month-wise Feat. Imp.'

    monthwise_df.to_excel(
        writer,
        sheet_name=ws_name,
        index=False
    )

    ws = writer.book[ws_name]

    style_header(
        ws,
        1
    )

    add_number_format(
        ws,
        2,
        ws.max_row
    )


# ------------------------------------------------------------
# Final workbook formatting
# ------------------------------------------------------------

wb = load_workbook(
    OUTPUT_EXCEL
)

for ws in wb.worksheets:

    ws.freeze_panes = 'A2'

    auto_width(
        ws
    )

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical='center',
                wrap_text=True
            )

wb.save(
    OUTPUT_EXCEL
)

print(
    '\nExcel workbook successfully generated:'
)

print(
    OUTPUT_EXCEL
)


# ============================================================
# 17. FINAL SUMMARY
# ============================================================

print('\n' + '=' * 70)
print('PIPELINE COMPLETED')
print('=' * 70)

print(
    '\nOriginal dataset:',
    data.shape
)

print(
    'Correlation-reduced dataset:',
    data2a.shape
)

print(
    'Final retained original dataset:',
    Xy0_merge.shape
)

print(
    'Final rejected original dataset:',
    Xy0_merge_r.shape
)

print(
    'Final retained model X:',
    X_merge.shape
)

print(
    'Final retained model y:',
    y_merge.shape
)

print(
    '\nFinal 5-fold mean R²:',
    cv_results_df['R2'].mean()
)

print(
    'Final 5-fold SD R²:',
    cv_results_df['R2'].std()
)

print(
    '\nExcel output:',
    OUTPUT_EXCEL
)

print(
    '\nSheets generated:'
)

for sheet in [
    'Model Performance',
    'Comp. Study',
    'Ablation',
    'Topographical Param.',
    'Reduced Feat. Imp.',
    'Reduced Month-wise Feat. Imp.'
]:

    print(
        ' -',
        sheet
    )
